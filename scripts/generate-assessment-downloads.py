#!/usr/bin/env python3
"""Build original, source-attributed RU professional aids into public/downloads/assessment.

Source content here is an educational template, NOT a diagnostic questionnaire or
clinically validated measurement tool. No proprietary questions/scoring are copied.
Requires: reportlab and python-docx. Run before Astro's static build.
"""
from __future__ import annotations

import csv
from html import escape
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

OUT = Path(__file__).resolve().parents[1] / "public" / "downloads" / "assessment"
OUT.mkdir(parents=True, exist_ok=True)
DATE = "14.10.2026"
VERSION = "1.0"
BLUE = colors.HexColor("#245786")
PALE = colors.HexColor("#EBF3FA")
INK = colors.HexColor("#1D344F")
MUTED = colors.HexColor("#596B7B")
BORDER = colors.HexColor("#CEDCE9")


def font_file(name):
    roots = [
        Path("/usr/share/fonts/truetype/dejavu"),
        Path("/usr/share/fonts/truetype/liberation2"),
        Path("/usr/share/fonts/truetype/liberation"),
        Path("/usr/local/share/fonts"),
    ]
    for root in roots:
        candidate = root / name
        if candidate.exists():
            return str(candidate)
    raise FileNotFoundError("Cyrillic TrueType font unavailable on runner: " + name)


pdfmetrics.registerFont(TTFont("Hub", font_file("DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("Hub-Bold", font_file("DejaVuSans-Bold.ttf")))
pdfmetrics.registerFontFamily("Hub", normal="Hub", bold="Hub-Bold")
style = getSampleStyleSheet()
style.add(ParagraphStyle(name="HubTitle", fontName="Hub-Bold", fontSize=16, leading=22,
                         textColor=BLUE, spaceAfter=12))
style.add(ParagraphStyle(name="HubH2", fontName="Hub-Bold", fontSize=11, leading=17,
                         textColor=BLUE, spaceBefore=12, spaceAfter=7))
style.add(ParagraphStyle(name="HubBody", fontName="Hub", fontSize=8.8, leading=14,
                         textColor=INK, spaceAfter=8))
style.add(ParagraphStyle(name="HubSmall", fontName="Hub", fontSize=7.3, leading=11,
                         textColor=MUTED, spaceAfter=7))
style.add(ParagraphStyle(name="HubCell", fontName="Hub", fontSize=7.4, leading=11,
                         textColor=INK))
style.add(ParagraphStyle(name="HubHeadCell", fontName="Hub-Bold", fontSize=7.4, leading=11,
                         textColor=INK))


def P(text, name="HubBody"):
    return Paragraph(escape(text), style[name])


def footer(canvas, doc):
    canvas.saveState()
    width, height = doc.pagesize
    canvas.setStrokeColor(BORDER)
    canvas.line(18*mm, 18*mm, width-18*mm, 18*mm)
    canvas.setFont("Hub", 7)
    canvas.setFillColor(MUTED)
    canvas.drawString(18*mm, 13*mm, "Autism Hub Kazakhstan | редакционный инструмент, не диагностический тест")
    canvas.drawRightString(width-18*mm, 13*mm, f"v{VERSION} | {DATE} | {doc.page}")
    canvas.restoreState()


def pdf(name, blocks, wide=False):
    path=OUT / name
    doc=SimpleDocTemplate(str(path), pagesize=landscape(A4) if wide else A4,
                          leftMargin=18*mm, rightMargin=18*mm,
                          topMargin=16*mm, bottomMargin=24*mm,
                          title=name, author="Autism Hub Kazakhstan")
    doc.build(blocks, onFirstPage=footer, onLaterPages=footer)
    return path


def section(title, paragraphs):
    blocks = [P(title,"HubH2")]
    blocks.extend(P(x) for x in paragraphs)
    return blocks


def table(rows, col_widths, header=True):
    converted=[]
    for idx,row in enumerate(rows):
        converted.append([P(str(val), "HubHeadCell" if idx==0 and header else "HubCell") for val in row])
    t=Table(converted, colWidths=col_widths, repeatRows=1 if header else 0, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),PALE if header else colors.white),
        ("GRID",(0,0),(-1,-1),0.45,BORDER),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("LEFTPADDING",(0,0),(-1,-1),8), ("RIGHTPADDING",(0,0),(-1,-1),8),
        ("TOPPADDING",(0,0),(-1,-1),7), ("BOTTOMPADDING",(0,0),(-1,-1),7),
    ]))
    return t


def first():
    blocks=[
        P("Маршрут раннего выявления и направления", "HubTitle"),
        P("Профессиональная схема | Autism Hub Kazakhstan | версия 1.0, 14.10.2026", "HubSmall"),
        P("Схема помогает организовать разговор, оценку и направление. Это не клинический протокол Республики Казахстан и не диагностический тест."),
    ]
    phases=[
        ("1. Обнаружено опасение","Сообщения семьи / самого ребёнка, наблюдения врача или педагога, изменения в развитии."),
        ("2. Уточнён контекст","История развития и здоровья, язык общения, конкретные наблюдения в разных условиях."),
        ("3. Исключены причины для срочных действий","Регресс навыков, судороги, острые симптомы и иные медицинские признаки требуют соответствующей своевременной оценки."),
        ("4. Решён вопрос о скрининге","Если показан: выбрать подходящий возрасту и языку инструмент, применить авторские инструкции, при необходимости Follow-Up."),
        ("5. Объяснён результат","Положительный скрининг не подтверждает РАС. Отрицательный не исключает РАС при обоснованных опасениях."),
        ("6. Согласован профессиональный маршрут","Направить к компетентному клиницисту/службе при показаниях. Параллельно обсуждать помощь по потребностям."),
        ("7. Составлен понятный план","Зафиксировать следующий шаг, контакт, способ повторной связи, основания передачи данных и участие семьи."),
    ]
    for title,body in phases:
        blocks.append(table([[title],[body]],[174*mm],header=False))
        blocks.append(Spacer(1,2.5*mm))
    blocks.extend(section("Красные флаги",[
      "Новая утрата ранее освоенных социальных, речевых или двигательных навыков, судорожные приступы, острое ухудшение или выраженные медицинские симптомы требуют медицинской оценки по характеру риска. При непосредственной угрозе действуют неотложные службы.",
    ]))
    blocks.extend(section("Обязательные ограничения",[
      "Педагогическое наблюдение, скрининг и оценка ПМПК не равны установлению медицинского диагноза. Показания к дополнительным обследованиям устанавливаются врачом. Порядок направления и согласия проверяется по действующим правилам Казахстана.",
    ]))
    blocks.append(P("Источники: NICE CG128, https://www.nice.org.uk/guidance/cg128/chapter/recommendations ; CPS 2019, https://cps.ca/en/documents/position/asd-early-detection ; Кодекс РК https://adilet.zan.kz/rus/docs/K2000000360", "HubSmall"))
    return pdf("early-detection-pathway-ru.pdf",blocks)


COMPARE=[
  {
    "tool":"M-CHAT-R/F","kind":"Двухэтапный скрининг (ответы родителей + Follow-Up)",
    "age":"16-30 месяцев (авторы)",
    "scope":"Выявляет основания для дальнейшей оценки; не ставит диагноз",
    "requirements":"Следовать оригинальной процедуре и условиям использования. Проверить язык и адаптацию. Не размещать копию на сайте без лицензии.",
    "url":"https://www.mchatscreen.com/mchat-rf/",
  },
  {
    "tool":"ADOS-2","kind":"Стандартизированное диагностическое наблюдение",
    "age":"Модуль по языку и развитию; применимость по руководству",
    "scope":"Дополняет комплексную клиническую оценку; не является автономным диагнозом",
    "requirements":"Лицензированные материалы, профессиональная подготовка, проверка валидности перевода и применимых условий.",
    "url":"https://www.wpspublish.com/ados-2-autism-diagnostic-observation-schedule-second-edition",
  },
  {
    "tool":"ADI-R","kind":"Стандартизированное интервью с информантом об истории развития",
    "age":"Возраст и уровень развития по официальному руководству",
    "scope":"Дополняет клинические сведения, не является универсальным скринингом",
    "requirements":"Право использования и обучения, подходящий информант, языковая валидность, клиническая интерпретация.",
    "url":"https://www.wpspublish.com/adi-r-autism-diagnostic-interview-revised",
  },
  {
    "tool":"CARS-2","kind":"Рейтинговая оценка по наблюдениям и клиническим сведениям",
    "age":"С 2 лет; варианты по характеристикам ребёнка (издатель)",
    "scope":"Помогает структурировать оценку; не заменяет клинические критерии",
    "requirements":"Материалы Pearson, указанная издателем квалификация, версия и возрастные ограничения.",
    "url":"https://www.pearsonassessments.com/en-us/en-us/Store/Professional-Assessments/Behavior/Childhood-Autism-Rating-Scale-%7C-Second-Edition/p/100000164",
  },
]


def comparison():
    out=OUT/"diagnostic-instrument-comparison-ru.csv"
    with out.open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f,delimiter=";")
        writer.writerow(["Инструмент","Тип","Возраст / группа","Назначение","Условия и ограничения","Первоисточник","Версия","Статус"])
        for v in COMPARE:
            writer.writerow([v["tool"],v["kind"],v["age"],v["scope"],v["requirements"],v["url"],"1.0 / 14.10.2026","Редакционная таблица, не диагностика"])
    blocks=[
      P("Диагностические инструменты: сравнительная таблица","HubTitle"),
      P("Редакционная справка AutismHub | без бланков, тестовых вопросов, инструкций подсчёта и нормативов", "HubSmall"),
      P("Точная применимость, обучение и право использования определяются по руководству правообладателя. Диагноз не устанавливается по одному инструменту."),
    ]
    rows=[["Инструмент","Вид / задача","Возраст и условия","Главное ограничение"]]
    for v in COMPARE:
        rows.append([v["tool"],v["kind"],v["age"],v["scope"]+"; "+v["requirements"]])
    blocks.append(table(rows,[37*mm,70*mm,57*mm,95*mm]))
    blocks.append(Spacer(1,7*mm))
    blocks.append(P("Официальные сайты: M-CHAT-R/F - mchatscreen.com/mchat-rf/; ADOS-2 и ADI-R - wpspublish.com; CARS-2 - pearsonassessments.com. NICE CG128: nice.org.uk/guidance/cg128/chapter/recommendations", "HubSmall"))
    blocks.append(P("Для редактирования и фильтрации используйте CSV в комплекте; не загружайте в него персональные сведения детей в общедоступных системах.", "HubSmall"))
    pdf("diagnostic-instrument-comparison-ru.pdf", blocks, wide=True)
    return out


FIELDS=[
 ("Причина обращения", ["Что беспокоит ребёнка и/или семью?", "Как давно это наблюдается?", "Какие цели у консультации?"]),
 ("Языковая и коммуникационная среда", ["Какие языки ребёнок понимает и использует?", "Как он выражает просьбу, отказ, боль?", "Нужны ли AAC или переводчик?"]),
 ("Развитие и ранняя история", ["Сведения о беременности, родах и раннем развитии по показаниям", "Когда появились особенности общения, игры и поведения?", "Были ли утраты уже сформированных навыков?"]),
 ("Состояние здоровья", ["Слух, зрение, эпизоды судорог и иные симптомы", "Сон, питание, боль, желудочно-кишечные жалобы", "Предыдущие обследования и значимые медицинские сведения"]),
 ("Повседневное участие", ["Участие дома, в саду/школе, общение со сверстниками", "Какие условия помогают или затрудняют участие?", "Что ребёнок умеет и любит делать?"]),
 ("Другие источники", ["Наблюдения педагогов и иных специалистов", "Имеющиеся документы, инструменты, языки оценки", "Что нужно уточнить или запросить дополнительно?"]),
 ("Следующие шаги", ["Какие вопросы остаются без ответа?", "Нужна ли консультация профильного врача или специалиста?", "Какая поддержка возможна на основании выявленных потребностей?", "Кто ответственен и когда семья сможет обратиться повторно?"]),
]


def history_pdf():
    blocks=[
      P("Бланк истории развития и интервью с семьёй","HubTitle"),
      P("Редакционный шаблон AutismHub | версия 1.0 · 14.10.2026 | не диагностическая шкала","HubSmall"),
      P("Шаблон для уполномоченного специалиста. Нет баллов, клинических порогов или диагностического результата. Заполняйте только в разрешённой системе организации; не загружайте данные ребёнка на сайт."),
    ]
    for ix,(title, questions) in enumerate(FIELDS,1):
        blocks.append(P(f"{ix}. {title}","HubH2"))
        for q in questions:
            blocks.append(P("• "+q,"HubBody"))
        blocks.append(P("Наблюдения, слова семьи и ограничения данных: ____________________________________________","HubSmall"))
        blocks.append(P("________________________________________________________________________________","HubSmall"))
    blocks.extend(section("Профессиональные ограничения",[
      "Это не медицинское заключение, не форма информированного согласия и не анкета с нормативными баллами. Используйте актуальные утверждённые формы и правила Казахстана; объём медицинских данных и законность передачи определяются отдельно.",
    ]))
    blocks.append(P("Основа структуры: NICE CG128, разделы 1.5 и 1.8; https://www.nice.org.uk/guidance/cg128/chapter/recommendations", "HubSmall"))
    return pdf("developmental-history-template-ru.pdf",blocks)


def history_docx():
    d=Document()
    sec=d.sections[0]
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.0);sec.right_margin=Cm(2.0)
    normal=d.styles["Normal"];normal.font.name="DejaVu Sans";normal.font.size=Pt(9)
    normal.paragraph_format.space_after=Pt(6)
    title=d.add_heading("Бланк истории развития и интервью с семьёй",0)
    for run in title.runs: run.font.name="DejaVu Sans";run.font.color.rgb=RGBColor(36,87,134)
    p=d.add_paragraph("Autism Hub Kazakhstan · версия 1.0 · 14.10.2026 · редактируемый профессиональный шаблон")
    p.style="Subtitle"
    d.add_paragraph("Не является диагностической шкалой, медицинским заключением или формой согласия. Используйте только в соответствии с требованиями учреждения и законом о защите данных.")
    for i,(heading, questions) in enumerate(FIELDS,1):
        d.add_heading(f"{i}. {heading}", level=2)
        t=d.add_table(rows=0,cols=2)
        t.style="Table Grid"
        for question in questions:
            cells=t.add_row().cells
            cells[0].text=question
            cells[1].text=" "
        d.add_paragraph("Дополнительные наблюдения / неопределённость данных: ____________________________________")
    d.add_heading("Основание и границы применения",level=2)
    d.add_paragraph("Составлено AutismHub по структуре профессиональной оценки NICE CG128. Вопросы авторские, без перепечатки защищённых диагностических инструментов.")
    d.add_paragraph("Источник: https://www.nice.org.uk/guidance/cg128/chapter/recommendations")
    d.add_paragraph("Право РК: https://adilet.zan.kz/rus/docs/K2000000360 и https://adilet.zan.kz/rus/docs/Z1300000094")
    for i in range(1,4):
        d.add_paragraph("")
    d.save(OUT/"developmental-history-template-ru.docx")
    return OUT/"developmental-history-template-ru.docx"


if __name__ == "__main__":
    outputs=[first(),comparison(),history_pdf(),history_docx()]
    expected={x.name for x in outputs}
    expected.add("diagnostic-instrument-comparison-ru.pdf")
    for filename in sorted(expected):
        p=OUT/filename
        assert p.exists() and p.stat().st_size>1000, f"Missing or suspicious download: {p}"
        print(f"ASSESSMENT_RESOURCE_OK {p.name} {p.stat().st_size}")
