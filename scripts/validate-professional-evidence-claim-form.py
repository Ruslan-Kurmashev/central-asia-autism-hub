#!/usr/bin/env python3
"""Structural validation of the original claims-appraisal PDF and DOCX."""
import zipfile
from pathlib import Path
from docx import Document
from pypdf import PdfReader
import sys
ROOT=Path(__file__).resolve().parents[1]
out=ROOT/"public"/"downloads"/"evidence"
pdf=out/"critical-claims-appraisal-ru.pdf"
docx=out/"critical-claims-appraisal-ru.docx"
assert pdf.exists() and docx.exists(),"Evidence appraisal files missing"
assert pdf.read_bytes().startswith(b"%PDF"),"Malformed printable PDF"
reader=PdfReader(str(pdf),strict=True)
assert 3<=len(reader.pages)<=12,("Unexpected PDF pagination",len(reader.pages))
all_text="\n".join(p.extract_text() or "" for p in reader.pages)
for keyword in ["Критическая оценка заявлений","PICO","безопасност","Не присваивайте","Cochrane","NICE"]:
    assert keyword in all_text,("Missing PDF content",keyword)
assert all(len(p.extract_text() or "")>70 for p in reader.pages),"Empty printable page"
with zipfile.ZipFile(docx) as z:
    assert z.testzip() is None,"Invalid editable DOCX ZIP"
document=Document(str(docx))
assert len(document.tables)==6,("Editable section count",len(document.tables))
assert sum(len(t.rows) for t in document.tables)>=30,"Insufficient editable fields"
assert len([p for p in document.paragraphs if p.style.name=="Heading 1"])>=7
assert "Критическая оценка" in "\n".join(p.text for p in document.paragraphs)
print("EVIDENCE_CLAIMS_FORM_VALID",len(reader.pages),"PDF pages",len(document.tables),"DOCX sections")
