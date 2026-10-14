#!/usr/bin/env python3
"""Reproducibly build AutismHub download files for static GitHub Pages deployment.

Do not accept patient inputs or publish filled medical records. Outputs are
generic educational professional aids and not diagnostic measures.
"""
from pathlib import Path
from io import BytesIO
import csv
import html
import zipfile
from xml.sax.saxutils import escape
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from docx import Document
from docx.shared import Cm, Pt, RGBColor

OUT=Path(__file__).resolve().parents[1]/"public"/"downloads"/"assessment"
OUT.mkdir(parents=True,exist_ok=True)
font="/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
bold="/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
if not Path(font).exists():
    font="/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"
    bold="/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"
pdfmetrics.registerFont(TTFont("Hub",font))
pdfmetrics.registerFont(TTFont("HubBold",bold))
INK=HexColor("#26384E");NAVY=HexColor("#1D4B79");PAPER=HexColor("#EBF3FA");BORDER=HexColor("#C8D7E5")

def paragraph(text,size=8,bold=False,color=INK):
    return Paragraph(escape(text),ParagraphStyle("p",fontName="HubBold" if bold else "Hub",fontSize=size,leading=size*1.45,textColor=color,spaceAfter=4))

def text(c,x,y,s,size=9,b=False,color=INK):
    c.setFont("HubBold" if b else "Hub",size);c.setFillColor(color);c.drawString(x,y,s)

def pathway():
    c=canvas.Canvas(str(OUT/"early-detection-pathway-ru.pdf"),pagesize=A4)
    w,h=A4;c.setTitle("AutismHub - Раннее выявление и направление")
    c.setFillColor(NAVY);c.rect(0,h-105,w,105,stroke=0,fill=1)
    text(c,38,h-39,"AUTISMHUB / ПРАКТИЧЕСКИЕ МАТЕРИАЛЫ",9,True,colors.white)
    text(c,38,h-70,"Раннее выявление: маршрут действий",16,True,colors.white)
    text(c,38,h-89,"Версия 1.0 · 14.10.2026 · Печатная схема",7,False,colors.white)
    steps=[
        ("Обнаружены признаки или запрос семьи",["Наблюдения, история развития, предпочтения ребёнка,","изменение навыков и вопросы родителей."]),
        ("Уточните клинический контекст",["Языки семьи, возраст, соматические жалобы; регресс","навыков и другие важные симптомы оценит врач."]),
        ("Наблюдение и скрининг по показаниям",["Только подходящий возрасту и языку инструмент.","Скрининг не ставит и не исключает диагноз РАС."]),
        ("Нужна ли дополнительная оценка?",["При положительном результате или устойчивых опасениях","рассмотрите профессиональное направление."]),
        ("Поддержка по актуальным потребностям",["Согласуйте доступную помощь коммуникации и участию.","Не обещайте недоступные услуги и сроки."]),
        ("Письменный итог и дальнейший контакт",["Кто отвечает, куда обращаться и когда пересмотреть","ситуацию; учитывайте правила согласия и данных."])
    ]
    margin=37;bw=w-74;boxh=75;y=h-122
    for i,(head,lines) in enumerate(steps):
        y-=boxh+ (13 if i else 0)
        c.setFillColor(PAPER if i%2==0 else colors.white);c.setStrokeColor(BORDER)
        c.roundRect(margin,y,bw,boxh,10,stroke=1,fill=1)
        c.setFillColor(HexColor("#196E76"));c.circle(58,y+boxh-24,14,stroke=0,fill=1)
        text(c,54,y+boxh-28,str(i+1),10,True,colors.white)
        text(c,83,y+boxh-25,head,9.5,True,NAVY)
        for j,line in enumerate(lines):text(c,83,y+boxh-43-j*12,line,7.3)
        if i<5:
            c.setStrokeColor(NAVY);c.line(w/2,y-1,w/2,y-12)
    c.setStrokeColor(HexColor("#E5CC90"));c.setFillColor(HexColor("#FFF8E9"));c.roundRect(37,55,w-74,80,8,stroke=1,fill=1)
    for j,line in enumerate(["Не диагностический алгоритм и не официальный протокол Казахстана.",
                              "При медицинских тревожных признаках - оценка врача по показаниям.",
                              "Скрининг и медицинский диагноз различаются, как и функции ПМПК.",
                              "Не вводите данные ребёнка в публичные сервисы."]):
        text(c,50,118-j*14,line,7.1)
    text(c,37,36,"Источники: NICE CG128; CPS Early detection 2019; Кодекс РК № 360-VI.",6.9)
    text(c,37,22,"nice.org.uk/guidance/cg128  |  cps.ca/en/documents/position/asd-early-detection",6.7)
    c.showPage();c.save()

HEAD=["Инструмент","Основное назначение","Возраст / применение","Требования","Ограничения","Условия использования","Официальный источник"]
ROWS=[
 ["M-CHAT-R/F","Скрининг признаков РАС","Дети 16–30 месяцев (правила авторов)","Анкета + Follow-Up по авторской процедуре","Скрининг, не диагноз","Проверить права и языковую версию","https://www.mchatscreen.com/"],
 ["ADOS-2","Стандартизированное клиническое наблюдение","Модули по языку и уровню развития","Обучение, комплект, профессиональная оценка","Не является самостоятельным диагнозом","Защищённые материалы, лицензия","https://www.wpspublish.com/"],
 ["ADI-R","Структурированное интервью с информантом","Возраст/развитие по оригинальному руководству","Компетентный интервьюер, анамнез развития","Зависит от качества воспоминаний","Защищённые материалы, лицензия","https://www.wpspublish.com/"],
 ["CARS-2","Клиническая рейтинговая шкала","Версия и возраст по руководству","Подготовка, многокомпонентная оценка","Одна шкала не заменяет клиническую оценку","Защищённые материалы, лицензия","https://www.wpspublish.com/"],
 ["ASQ-3","Общий скрининг развития","Возрастные формы по руководству","Языковая и возрастная версия","Не диагностирует аутизм","Издательские условия","https://agesandstages.com/"],
]
DISCLAIMER="Ресурс AutismHub не является шкалой. Проверяйте актуальные возрастные рамки, языковую валидность, лицензию и обучение по инструкциям правообладателей."

def comparisons():
    with (OUT/"diagnostic-instrument-comparison-ru.csv").open("w",encoding="utf-8-sig",newline="") as handle:
        w=csv.writer(handle,delimiter=";");w.writerow(HEAD);w.writerows(ROWS)
    doc=SimpleDocTemplate(str(OUT/"diagnostic-instrument-comparison-ru.pdf"),pagesize=landscape(A4),
       leftMargin=17*mm,rightMargin=17*mm,topMargin=15*mm,bottomMargin=16*mm)
    story=[paragraph("AUTISMHUB / СРАВНИТЕЛЬНАЯ ТАБЛИЦА",8,True,NAVY),
           paragraph("Инструменты оценки развития и аутизма",15,True,NAVY),
           paragraph("Версия 1.0 · 14.10.2026 · Профессиональная справка, не клинический тест.")]
    widths=[68,106,142,143,157,119]
    hd=[paragraph(x,6.3,True,colors.white) for x in HEAD[:6]]
    data=[hd]+[[paragraph(x,6.7) for x in row[:6]] for row in ROWS]
    t=Table(data,colWidths=widths,repeatRows=1,hAlign="LEFT")
    t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),NAVY),("VALIGN",(0,0),(-1,-1),"TOP"),
      ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,PAPER]),("GRID",(0,0),(-1,-1),.5,BORDER),
      ("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),6),
      ("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8)]))
    story.extend([Spacer(1,12),t,Spacer(1,14),paragraph(DISCLAIMER,8,True),
      paragraph("Первичные источники: mchatscreen.com; wpspublish.com; agesandstages.com. Ни один инструмент не заменяет полноценную диагностику.",7)])
    doc.build(story)

def history():
    d=Document();s=d.sections[0]
    s.top_margin=Cm(1.7);s.bottom_margin=Cm(1.6);s.left_margin=Cm(2);s.right_margin=Cm(2)
    normal=d.styles["Normal"];normal.font.name="DejaVu Sans";normal.font.size=Pt(9);normal.font.color.rgb=RGBColor(37,55,77)
    for sty,size in [("Title",17),("Heading 1",12)]:
        obj=d.styles[sty];obj.font.name="DejaVu Sans";obj.font.size=Pt(size)
        obj.font.bold=True;obj.font.color.rgb=RGBColor(27,75,120)
    s.header.paragraphs[0].text="AUTISMHUB / ПРОФЕССИОНАЛЬНЫЕ РАБОЧИЕ МАТЕРИАЛЫ"
    s.footer.paragraphs[0].text="Учебный шаблон v1.0 · 14.10.2026 | Не размещать персональные данные в публичных сервисах"
    d.add_heading("История развития ребёнка",0)
    d.add_paragraph("Редактируемый шаблон интервью для компетентных специалистов. Не валидированная диагностическая шкала, не медицинский протокол или обязательный бланк Республики Казахстан.")
    d.add_paragraph("Сведения о ребёнке следует хранить исключительно в одобренной учреждением защищённой системе.")
    def block(name,note,n=3):
        d.add_heading(name,1);d.add_paragraph(note)
        for _ in range(n):d.add_paragraph("_______________________________________________________________________________________")
    block("A. Обращение и условия","Дата консультации, возраст ребёнка, профессиональная роль сотрудника; причина обращения семьи.",2)
    block("Предпочтительные языки и коммуникация","Разделяйте язык семьи, языки ребёнка, способы общения, потребность в переводчике.",2)
    block("Сильные стороны и интересы","Как ребёнок выражает пожелания и отказы, что помогает участию.",2)
    d.add_page_break()
    block("B. Раннее развитие и изменения навыков","Факты истории развития, используемые клинические документы, замеченные изменения.",3)
    block("Социальная коммуникация","Примеры общения и игры дома, в саду и других знакомых условиях.",3)
    block("Языки, речь и альтернативная коммуникация","Понимание, жесты, изображения, AAC и особенности каждого языка.",3)
    block("Поведение, интересы, сенсорные особенности","Опишите конкретную ситуацию, среду и реакцию; не приписывайте мотивы.",2)
    d.add_page_break()
    block("C. Здоровье и повседневная жизнь","Сон, питание, слух, зрение, другие жалобы и показания к профильной оценке.",3)
    block("Сведения из других ситуаций","Наблюдения семьи и специалистов с указанием источника и контекста.",2)
    block("Недостающие данные и диагностическая неопределённость","Какие сведения необходимы и кто уполномочен проводить дополнительную оценку.",2)
    block("Согласованный следующий шаг","Ответственный, сроки связи, направления и правила передачи информации.",2)
    d.add_heading("Источники и ограничения",1)
    for item in ["NICE CG128: https://www.nice.org.uk/guidance/cg128/chapter/Recommendations",
      "Canadian Paediatric Society: https://cps.ca/en/documents/position/asd-diagnostic-assessment",
      "Кодекс РК: https://adilet.zan.kz/rus/docs/K2000000360",
      "Шаблон AutismHub не воспроизводит ADI-R, ADOS-2, CARS-2 и другие защищённые инструменты."]:
        d.add_paragraph(item)
    file=OUT/"developmental-history-template-ru.docx";d.save(file)
    # PDF printable version from DOCX via LibreOffice, which is installed on the build runner.
    import subprocess
    subprocess.run(["libreoffice","-env:UserInstallation=file:///tmp/autismhub-lo-profile",
      "--headless","--convert-to","pdf","--outdir",str(OUT),str(file)],check=True,
      stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=90)
    assert (OUT/"developmental-history-template-ru.pdf").exists()

def spreadsheet():
    # XLSX via the ISO OpenXML Spreadsheet package. No macro, code or external links.
    cols="ABCDEFG"
    data=[["AutismHub | Инструменты оценки развития и РАС"],["Версия 1.0 · 14.10.2026 | Источники и лицензии сверять отдельно."],[],HEAD,*ROWS,[],[DISCLAIMER]]
    def xmltext(x):return escape(str(x),{'"':"&quot;"})
    rows=[]
    for i,line in enumerate(data,1):
        cells="".join(f'<c r="{cols[j]}{i}" t="inlineStr"><is><t>{xmltext(v)}</t></is></c>' for j,v in enumerate(line))
        rows.append(f'<row r="{i}">{cells}</row>')
    widths=[18,32,40,41,44,37,40]
    coltag="".join(f'<col min="{i+1}" max="{i+1}" width="{w}" customWidth="1"/>' for i,w in enumerate(widths))
    sheet=f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><cols>{coltag}</cols><sheetData>{"".join(rows)}</sheetData></worksheet>'
    types='''<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>'''
    rels='''<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>'''
    book='''<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Инструменты" sheetId="1" r:id="rId1"/></sheets></workbook>'''
    bookrels='''<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>'''
    with zipfile.ZipFile(OUT/"diagnostic-instrument-comparison-ru.xlsx","w",zipfile.ZIP_DEFLATED) as z:
      for name,contents in [("[Content_Types].xml",types),("_rels/.rels",rels),("xl/workbook.xml",book),("xl/_rels/workbook.xml.rels",bookrels),("xl/worksheets/sheet1.xml",sheet)]:
        z.writestr(name,contents.encode())

if __name__=="__main__":
    pathway();comparisons();history();spreadsheet()
    for file in sorted(OUT.iterdir()):
        if file.is_file():
            assert file.stat().st_size>1000
            print(file.name,file.stat().st_size)
