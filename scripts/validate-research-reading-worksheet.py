#!/usr/bin/env python3
"""Validate the printable/editable original scientific-paper reading worksheet."""
from pathlib import Path
import zipfile
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "public" / "downloads" / "research"
PDF = OUT / "scientific-article-reading-ru.pdf"
DOCX = OUT / "scientific-article-reading-ru.docx"

def main():
    assert PDF.read_bytes()[:4] == b"%PDF"
    pdf = PdfReader(str(PDF), strict=True)
    assert 1 <= len(pdf.pages) <= 10, "Unexpected PDF pagination"
    text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    for phrase in ["Читаем научную статью", "ПЕРВОИСТОЧНИКИ", "ПРАВИЛА ИСПОЛЬЗОВАНИЯ"]:
        assert phrase in text, phrase
    assert all(len(page.extract_text() or "") > 90 for page in pdf.pages)
    with zipfile.ZipFile(DOCX) as z:
        assert z.testzip() is None
    doc = Document(str(DOCX))
    assert len(doc.tables) == 4, "Expected four editable sections"
    assert all(len(table.rows) == 4 and len(table.columns) == 2 for table in doc.tables)
    assert any("Читаем научную статью" in p.text for p in doc.paragraphs)
    for f in [PDF, DOCX]:
        assert f.stat().st_size > 2500
        print("SCIENCE_WORKSHEET_VALID", f.name, f.stat().st_size)
    print("SCIENCE_WORKSHEET_VALIDATED", len(pdf.pages), "PDF pages", len(doc.tables), "DOCX sections")

if __name__ == "__main__":
    main()
