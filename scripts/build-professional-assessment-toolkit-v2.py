#!/usr/bin/env python3
"""Build six original, unfilled professional assessment workbooks for AutismHub.

Source of truth: scripts/assessment-practical-tool-specs.json.
PDF (print) + DOCX (editable) for each resource, plus a two-week sleep
observation CSV. Do not copy protected diagnostic instruments or insert
any patient data. These are non-scored practice aids, not diagnostic tests.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table,
    TableStyle
)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "public" / "downloads" / "assessment"
OUT.mkdir(parents=True, exist_ok=True)
SPECS = ROOT / "scripts" / "assessment-practical-tool-specs.json"
DATE = "14.10.2026"
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
if not Path(FONT).exists():
    FONT = "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"
    BOLD = "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"
pdfmetrics.registerFont(TTFont("Hub2", FONT))
pdfmetrics.registerFont(TTFont("Hub2Bold", BOLD))

NAVY = HexColor("#1e4c78")
INK = HexColor("#26384e")
PALE = HexColor("#f1f6fb")
BORDER = HexColor("#c8dae8")
MUTED = HexColor("#64778b")
VERSION = "v1.0"

def p(text: str, size=8, *, strong=False, color=INK, after=0, leading=None):
    style = ParagraphStyle("tool", fontName="Hub2Bold" if strong else "Hub2",
                           fontSize=size, leading=leading or size * 1.45,
                           textColor=color, spaceAfter=after, splitLongWords=True)
    return Paragraph(escape(str(text)), style)

def pdf_footer(c, doc):
    width, _ = A4
    c.saveState()
    c.setStrokeColor(BORDER)
    c.line(17 * mm, 16 * mm, width - 17 * mm, 16 * mm)
    c.setFillColor(MUTED)
    c.setFont("Hub2", 7)
    c.drawString(17 * mm, 11 * mm, "AUTISMHUB · Авторский рабочий инструмент · " + DATE)
    c.drawRightString(width - 17 * mm, 11 * mm, str(doc.page))
    c.restoreState()

def build_pdf(spec):
    filename = spec["filename"] + "-ru.pdf"
    doc = SimpleDocTemplate(
        str(OUT / filename), pagesize=A4, topMargin=16 * mm,
        bottomMargin=21 * mm, leftMargin=17 * mm, rightMargin=17 * mm,
        title="AutismHub | " + spec["title"], author="AutismHub"
    )
    usable = A4[0] - 34 * mm
    left = 164
    right = usable - left
    story = [
        p("AUTISMHUB / РАБОЧИЙ МАТЕРИАЛ ДЛЯ СПЕЦИАЛИСТОВ", 8, strong=True, color=NAVY, after=7),
        p(spec["title"], 15, strong=True, color=NAVY, after=9),
        p("Версия " + VERSION + " · " + DATE + " · " + spec["audience"], 7.4, color=MUTED, after=8),
        p(spec["purpose"], 8.3, after=10),
        HRFlowable(width="100%", thickness=0.65, color=BORDER),
        Spacer(1, 10),
    ]
    for section in spec["sections"]:
        story.append(p(section["title"], 10.5, strong=True, color=NAVY, after=7))
        rows = []
        for label, hint in section["rows"]:
            lhs = [p(label, 8.2, strong=True), Spacer(1, 3), p(hint, 7, color=MUTED)]
            rhs = [Spacer(1, 37)]
            rows.append([lhs, rhs])
        table = Table(rows, colWidths=[left, right], hAlign="LEFT")
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, PALE]),
            ("LINEBELOW", (0, 0), (-1, -1), 0.45, BORDER),
            ("LINEBEFORE", (1, 0), (1, -1), 0.45, BORDER),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(table)
        story.append(Spacer(1, 15))
    story.append(p("ПРАВИЛА ИСПОЛЬЗОВАНИЯ", 9, strong=True, color=NAVY, after=7))
    for item in spec["notes"]:
        story.append(p("•  " + item, 7.6, after=4))
    story.append(p(
        "Не валидированная диагностическая шкала. Не обязательный бланк Республики Казахстан. "
        "Вносить и хранить реальные персональные данные можно только в разрешённых системах организации.",
        7.5, strong=True, after=8))
    story.append(p("ПЕРВОИСТОЧНИКИ (НЕ КОПИИ ТЕСТОВ)", 8.5, strong=True, color=NAVY, after=5))
    for label, url in spec["references"]:
        story.append(p(label + ": " + url, 7, color=MUTED, after=5))
    doc.build(story, onFirstPage=pdf_footer, onLaterPages=pdf_footer)
    return filename

def docx_styles(document):
    s = document.sections[0]
    s.top_margin = Cm(1.65); s.bottom_margin = Cm(1.55)
    s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)
    base = document.styles["Normal"]
    base.font.name = "DejaVu Sans"
    base.font.size = Pt(8.5)
    base.font.color.rgb = RGBColor(38, 56, 78)
    base.paragraph_format.space_after = Pt(4)
    for name, size in [("Title", 15), ("Heading 1", 11)]:
        style = document.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(30, 76, 120)
    header = s.header.paragraphs[0]
    header.text = "AUTISMHUB / ПРОФЕССИОНАЛЬНЫЕ РАБОЧИЕ МАТЕРИАЛЫ"
    header.style = document.styles["Normal"]
    footer = s.footer.paragraphs[0]
    footer.text = "Оригинальный бланк AutismHub · " + VERSION + " · " + DATE + " · Без данных пациентов"
    footer.style = document.styles["Normal"]

def build_docx(spec):
    filename = spec["filename"] + "-ru.docx"
    document = Document()
    docx_styles(document)
    document.add_heading(spec["title"], 0)
    intro = document.add_paragraph(spec["purpose"])
    intro.style = document.styles["Normal"]
    document.add_paragraph(
        "Версия " + VERSION + " · " + DATE + " · " + spec["audience"]
    )
    for section in spec["sections"]:
        document.add_heading(section["title"], 1)
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.columns[0].width = Cm(7.1)
        table.columns[1].width = Cm(9.1)
        for label, hint in section["rows"]:
            left, right = table.add_row().cells
            left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            lead = left.paragraphs[0]
            lead.add_run(label).bold = True
            help_text = left.add_paragraph(hint)
            help_text.style = document.styles["Normal"]
            space = right.paragraphs[0]
            space.add_run("_________________________________________________")
            # Placeholder-free blank area remains fully editable by a professional.
            right.add_paragraph(" ")
        document.add_paragraph(" ")
    document.add_heading("Ограничения использования", 1)
    for item in spec["notes"]:
        document.add_paragraph("• " + item)
    document.add_paragraph(
        "Это не валидированная диагностическая шкала, не официальный бланк Казахстана "
        "и не копия защищённой методики. Не передавайте заполненные документы в публичные сервисы."
    )
    document.add_heading("Первоисточники", 1)
    for label, url in spec["references"]:
        document.add_paragraph(label + ": " + url)
    document.save(OUT / filename)
    return filename

def sleep_csv():
    filename = "cooccurring-conditions-sleep-log-ru.csv"
    # NICE CG170 recommends a 2-week diary for sleep concerns. This table is
    # independently authored and deliberately does not collect names/IDs.
    headers = [
        "День", "Дата (заполнить локально)", "Начало подготовки ко сну",
        "Предположительное время засыпания", "Ночные пробуждения",
        "Время утреннего подъёма", "Дневной сон", "Храп / особенности дыхания",
        "Возможная боль или дискомфорт", "Примечания / условия"
    ]
    with (OUT / filename).open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(headers)
        for day in range(1, 15):
            writer.writerow([str(day), "", "", "", "", "", "", "", "", ""])
    return filename

def main():
    specs = json.loads(SPECS.read_text(encoding="utf-8"))
    assert len(specs) == 6
    assert len({x["filename"] for x in specs}) == 6
    created = []
    for spec in specs:
        assert len(spec["sections"]) >= 4
        created.extend([build_pdf(spec), build_docx(spec)])
    created.append(sleep_csv())
    for name in created:
        filepath = OUT / name
        assert filepath.stat().st_size > (90 if name.endswith(".csv") else 1300), name
        print("ASSESSMENT_RESOURCE_OK", name, filepath.stat().st_size)
    print("ASSESSMENT_TOOLKIT_V2_OK", len(created), "files", len(specs), "resource sets")

if __name__ == "__main__":
    main()
