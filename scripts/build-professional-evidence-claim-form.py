#!/usr/bin/env python3
"""Build the original AutismHub critical-claims appraisal form (PDF and DOCX).

Purpose: an unscored evidence/harms documentation aid, not a validated scale,
medical approval, regulatory assessment or a reproduction of another checklist.
Requires reportlab and python-docx; runs before Astro publication.
"""
from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from reportlab.lib import colors
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "public" / "downloads" / "evidence"
OUTPUT.mkdir(parents=True, exist_ok=True)
STEM = "critical-claims-appraisal-ru"
TITLE = "Критическая оценка заявлений о методе поддержки при аутизме"
VERSION = "1.0 / 14.10.2026"

SECTIONS = [
  ("01 / Паспорт проверяемого заявления", [
    ("Точная цитата обещания", "Запишите дословно, без перефразирования."),
    ("Источник и дата", "Адрес страницы, объявление, рекламный буклет или запись."),
    ("Поставщик и интересы", "Кто продаёт услугу? Финансирование, лицензии, авторы."),
    ("Целевая группа", "Возраст, потребности, условия применения."),
    ("Предлагаемые действия", "Процедура, интенсивность, длительность, квалификация."),
  ]),
  ("02 / Проверяемый вопрос PICO", [
    ("P / Для кого", "Участники, характеристики, ограничения применимости."),
    ("I / Что делают", "Конкретный протокол и условия исполнения."),
    ("C / Сравнение", "Обычная помощь, другая программа, отсутствие воздействия."),
    ("O / Обещанный результат", "Функциональный исход, способ измерения и срок."),
    ("Механизм отдельно", "Гипотеза о том, как метод работает; это не доказательство."),
  ]),
  ("03 / Источники и уверенность в выводе", [
    ("Дата и стратегия поиска", "Где, когда и по каким терминам искали доказательства."),
    ("Первичные исследования", "Название, ссылка/DOI, дизайн, выборка и исход."),
    ("Независимые обзоры и рекомендации", "Точная рекомендация, возраст и дата пересмотра."),
    ("Сопоставимость", "Участники и условия похожи на рассматриваемую ситуацию?"),
    ("Риск смещения", "Сравнение, слепая оценка, выбывание, отчётность, интересы."),
    ("Величина и точность результата", "Эффект, интервалы неопределённости, длительность."),
    ("Что не известно", "Нет исследований, отрицательные данные, непроверенные исходы."),
  ]),
  ("04 / Польза, риски и этические условия", [
    ("Значимость исхода для человека", "Коммуникация, доступность, участие, комфорт, автономия."),
    ("Физический и психологический риск", "Боль, стресс, сенсорная перегрузка, противопоказания."),
    ("Прямые и косвенные издержки", "Деньги, время, упущенная другая помощь, требования к семье."),
    ("Согласие и отказ", "Как выражаются предпочтения и сохраняется право прекратить?"),
    ("Данные о безопасности", "Систематический мониторинг или только заявления продавца?"),
    ("Статус в Казахстане", "Проверить показания, разрешения и профессиональные полномочия."),
  ]),
  ("05 / Критические сигналы в заявлении", [
    ("Гарантия или «излечение»", "Есть ли чрезмерные обещания для всех людей?"),
    ("Недоступность исходных данных", "Можно ли прочитать протокол, полный отчёт и ограничения?"),
    ("Давление и скрытые условия", "Навязывание оплаты, секретность, отговоры от второго мнения."),
    ("Перенос результатов", "Не приписывают ли результат одной программы всем методам?"),
  ]),
  ("06 / Итоговая формулировка — без общего балла", [
    ("Что действительно подтверждено", "Указать узкое утверждение и надёжные источники."),
    ("Главная неопределённость", "Что неизвестно и почему это важно для решения."),
    ("Риски и нерешённые вопросы", "Когда необходима дополнительная клиническая оценка."),
    ("Следующий шаг", "Какие документы запросить, что обсудить и когда пересмотреть."),
    ("Дата / ответственная роль", "Заполняйте служебно, без данных конкретного ребёнка."),
  ]),
]
SOURCES = [
 "Cochrane Handbook, Chapter 14: https://www.cochrane.org/authors/handbooks-and-manuals/handbook/current/chapter-14",
 "NICE CG170: https://www.nice.org.uk/guidance/cg170/chapter/recommendations",
 "NCCIH Autism: https://www.nccih.nih.gov/health/autism",
 "Acuña et al., AJOT 2025: https://doi.org/10.5014/ajot.2025.051023",
]
WARNINGS = [
 "Это авторский образовательный бланк, не валидированная шкала и не официальная форма Республики Казахстан.",
 "Не присваивайте методу суммарный балл «эффективности». Критический риск безопасности не компенсируется другими пунктами.",
 "Не размещайте заполненный документ с персональными данными ребёнка в открытом доступе или сторонних публичных сервисах.",
 "При угрозе здоровью необходима квалифицированная медицинская оценка, а не самостоятельное испытание метода.",
]
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
if not Path(FONT).exists():
    FONT = "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"
    BOLD = "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"
pdfmetrics.registerFont(TTFont("EH", FONT))
pdfmetrics.registerFont(TTFont("EHB", BOLD))
NAVY = HexColor("#205b8f")
INK = HexColor("#26384e")
GRAY = HexColor("#526b7c")
LINE = HexColor("#b9d2e4")
PALE = HexColor("#f0f6fb")

def paragraph(text: str, *, size=8, bold=False, color=INK, after=0, leading=None):
    style = ParagraphStyle(
        "evidence", fontName="EHB" if bold else "EH",
        fontSize=size, leading=leading or size*1.4, textColor=color,
        spaceAfter=after, splitLongWords=True, wordWrap="CJK")
    return Paragraph(escape(str(text)), style)

def footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(LINE)
    canvas.line(17*mm, 16*mm, A4[0]-17*mm, 16*mm)
    canvas.setFont("EH", 6.8)
    canvas.setFillColor(GRAY)
    canvas.drawString(17*mm, 10.5*mm, "AUTISMHUB / авторский бланк / 14.10.2026")
    canvas.drawRightString(A4[0]-17*mm, 10.5*mm, str(doc.page))
    canvas.restoreState()

def build_pdf():
    filename = OUTPUT / (STEM + ".pdf")
    doc = SimpleDocTemplate(
        str(filename), pagesize=A4, topMargin=18*mm,
        bottomMargin=22*mm, leftMargin=17*mm, rightMargin=17*mm,
        title="AutismHub | " + TITLE, author="AutismHub",
    )
    width = A4[0] - 34*mm
    story = [
        paragraph("AUTISMHUB / ИНСТРУМЕНТ ДЛЯ СПЕЦИАЛИСТОВ",size=8,bold=True,color=NAVY,after=7),
        paragraph(TITLE,size=14.5,bold=True,color=NAVY,after=8),
        paragraph(VERSION + "  •  Один бланк для одного конкретного заявления",size=7.5,color=GRAY,after=9),
        paragraph("Задача: сформулировать обещание, проверить доказательства по конкретному исходу, документировать безопасность и неопределённость. Бланк не диагностирует и не сертифицирует метод.",size=8.2,after=10),
        HRFlowable(width="100%",thickness=0.75,color=LINE),
        Spacer(1,9),
    ]
    for heading, rows in SECTIONS:
        style = ParagraphStyle("section",fontName="EHB",fontSize=10.1,
                 leading=14,textColor=NAVY,spaceBefore=10,spaceAfter=7,keepWithNext=True)
        story.append(Paragraph(escape(heading), style))
        items=[]
        for name,hint in rows:
            label=[paragraph(name,bold=True,size=8.2,after=4),paragraph(hint,size=7,color=GRAY)]
            response=[Spacer(1,42)]
            items.append([label,response])
        table=Table(items,colWidths=[width*0.49,width*0.51],repeatRows=0,hAlign="LEFT")
        table.setStyle(TableStyle([
            ("VALIGN",(0,0),(-1,-1),"TOP"),
            ("ROWBACKGROUNDS",(0,0),(-1,-1),[colors.white,PALE]),
            ("LINEBELOW",(0,0),(-1,-1),0.5,LINE),
            ("LINEBEFORE",(1,0),(1,-1),0.5,LINE),
            ("TOPPADDING",(0,0),(-1,-1),8),
            ("BOTTOMPADDING",(0,0),(-1,-1),8),
            ("LEFTPADDING",(0,0),(-1,-1),9),
            ("RIGHTPADDING",(0,0),(-1,-1),9),
        ]))
        story.extend([table,Spacer(1,8)])
    story.append(paragraph("ОГРАНИЧЕНИЯ И БЕЗОПАСНОСТЬ",size=9.5,bold=True,color=NAVY,after=6))
    for line in WARNINGS:
        story.append(paragraph("• "+line,size=7.5,after=5))
    story.append(paragraph("ПЕРВОИСТОЧНИКИ ДЛЯ ПРОВЕРКИ",size=8.7,bold=True,color=NAVY,after=5))
    for source in SOURCES:
        story.append(paragraph(source,size=7,after=6,color=GRAY))
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return filename

def build_docx():
    filename = OUTPUT / (STEM + ".docx")
    document = Document()
    sec = document.sections[0]
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7)
    sec.left_margin=Cm(1.75);sec.right_margin=Cm(1.75)
    normal=document.styles["Normal"]
    normal.font.name="DejaVu Sans";normal.font.size=Pt(8.5)
    normal.font.color.rgb=RGBColor(38,56,78)
    normal.paragraph_format.space_after=Pt(5)
    for name,size in [("Title",15),("Heading 1",11)]:
        style=document.styles[name]
        style.font.name="DejaVu Sans";style.font.size=Pt(size)
        style.font.bold=True;style.font.color.rgb=RGBColor(32,91,143)
    sec.header.paragraphs[0].text="AUTISMHUB / ПРАКТИЧЕСКАЯ БИБЛИОТЕКА СПЕЦИАЛИСТА"
    sec.footer.paragraphs[0].text="Бланк 1.0 / 14.10.2026 / Не валидированная шкала / Не вводите персональные данные"
    document.add_heading(TITLE,0)
    document.add_paragraph(VERSION+" • Заполняйте один экземпляр для одного конкретного заявления.")
    document.add_paragraph("Фиксируйте первоисточники, неопределённость и безопасность. Документ не является клиническим заключением или регуляторным разрешением.")
    for heading, rows in SECTIONS:
        document.add_heading(heading,1)
        table=document.add_table(rows=0,cols=2)
        table.style="Table Grid"
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        table.columns[0].width=Cm(7.4)
        table.columns[1].width=Cm(9.2)
        for name,hint in rows:
            lhs,rhs=table.add_row().cells
            lhs.width=Cm(7.4);rhs.width=Cm(9.2)
            lhs.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            lead=lhs.paragraphs[0].add_run(name)
            lead.bold=True
            lhs.add_paragraph(hint)
            rhs.text=" "
            rhs.add_paragraph(" ")
        document.add_paragraph("")
    document.add_heading("Ограничения использования",1)
    for line in WARNINGS:
        document.add_paragraph("• "+line)
    document.add_heading("Первоисточники",1)
    for source in SOURCES:
        document.add_paragraph(source)
    document.save(filename)
    return filename

if __name__=="__main__":
    for item in [build_pdf(),build_docx()]:
        assert item.exists() and item.stat().st_size>3000
        print("EVIDENCE_CLAIMS_FORM_OK",item.name,item.stat().st_size)
