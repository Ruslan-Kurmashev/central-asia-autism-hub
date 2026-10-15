#!/usr/bin/env python3
"""Quality gate for five original printable and editable practitioner forms."""
from __future__ import annotations

import json
import zipfile
from pathlib import Path
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
SPECS = ROOT / "scripts" / "professional-practice-tool-specs.json"
OUT = ROOT / "public" / "downloads" / "practice"

def main():
    specs = json.loads(SPECS.read_text(encoding="utf-8"))
    assert len(specs) == 5
    for item in specs:
        name = item["filename"] + "-ru"
        pdf = OUT / (name + ".pdf")
        docx = OUT / (name + ".docx")
        assert pdf.read_bytes()[:4] == b"%PDF", "Invalid PDF: " + name
        reader = PdfReader(str(pdf), strict=True)
        assert 2 <= len(reader.pages) <= 15, (name, "wrong page count")
        full_text = "\n".join(page.extract_text() or "" for page in reader.pages)
        assert item["title"] in full_text, (name, "title missing from PDF")
        assert "ПРАВИЛА ИСПОЛЬЗОВАНИЯ" in full_text, (name, "limit missing from PDF")
        assert "ПЕРВОИСТОЧНИКИ" in full_text, (name, "references missing from PDF")
        assert all(len(page.extract_text() or "") > 70 for page in reader.pages), (name, "blank PDF page")
        with zipfile.ZipFile(docx) as z:
            assert z.testzip() is None, name + " DOCX ZIP corrupt"
        d = Document(str(docx))
        assert len(d.tables) == len(item["sections"]), (name, "wrong editable section count")
        for section, table in zip(item["sections"], d.tables):
            assert len(table.rows) == len(section["rows"]), (name, "missing editable rows")
            assert len(table.columns) == 2
        assert item["title"] in "\n".join(p.text for p in d.paragraphs)
        print("PRACTICE_FORM_VALID", name, len(reader.pages), "PDF pages", len(d.tables), "DOCX tables")
    print("PRACTICE_FORMS_VALID 5 resource sets / 10 files")

if __name__ == "__main__":
    main()
