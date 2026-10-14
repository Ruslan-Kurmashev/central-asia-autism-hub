#!/usr/bin/env python3
"""Check structure, pagination, language and source presence for all new assessment forms."""
import csv
import json
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"public"/"downloads"/"assessment"
specs=json.loads((ROOT/"scripts"/"assessment-practical-tool-specs.json").read_text(encoding="utf-8"))
assert len(specs)==6, "Six current sets expected"
for spec in specs:
    stem=spec["filename"]
    pdf=OUT/(stem+"-ru.pdf")
    docx=OUT/(stem+"-ru.docx")
    assert pdf.exists() and docx.exists(), stem+": missing generated resources"
    reader=PdfReader(str(pdf),strict=True)
    assert 2<=len(reader.pages)<=15,(stem,"unexpected page count",len(reader.pages))
    texts=[page.extract_text() or "" for page in reader.pages]
    combined="\n".join(texts)
    assert len(combined)>1200,(stem,"missing searchable PDF text")
    assert spec["title"] in combined,(stem,"title absent from PDF")
    assert all(len(page)>90 for page in texts),(stem,"empty PDF page")
    assert "Не валидированная" in combined,(stem,"clinical limitations omitted in PDF")
    with zipfile.ZipFile(docx) as z:
        assert z.testzip() is None,(stem,"damaged DOCX ZIP")
    doc=Document(str(docx))
    all_text="\n".join(p.text for p in doc.paragraphs)
    assert spec["title"] in all_text,(stem,"title absent from DOCX")
    assert len(doc.tables)==len(spec["sections"]),(stem,"missing editable sections")
    for table,section in zip(doc.tables,spec["sections"]):
        assert len(table.rows)==len(section["rows"]),(stem,"missing editable fields")
    assert all(any(link in combined for _,link in spec["references"][:1]) for _ in [0]),(stem,"source URLs missing in printable PDF")
    print("PRACTICE_TOOL_VALID",stem,"pdf_pages",len(reader.pages),"docx_tables",len(doc.tables))

with (OUT/"cooccurring-conditions-sleep-log-ru.csv").open(encoding="utf-8-sig",newline="") as f:
    rows=list(csv.reader(f,delimiter=";"))
assert len(rows)==15,"Sleep diary: header + 14 days required"
assert "Ночные пробуждения" in rows[0]
assert [row[0] for row in rows[1:]]==[str(n) for n in range(1,15)]
print("PRACTICE_TOOL_VALID sleep_diary 14 days")
print("PRACTICE_TOOL_VALIDATION_PASS",len(specs),"sets and 13 generated files")
